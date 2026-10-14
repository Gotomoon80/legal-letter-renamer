# -*- coding: utf-8 -*-
"""对照表导出：把重命名计划导出为 Excel，便于核对。"""
import os
import time
from typing import List

from . import config
from .logger import get_logger
from .rename_engine import PlanRow, STATUS_TEXT

log = get_logger("report")

_HEADERS = ["序号", "原文件名", "识别序号", "匹配姓名", "新文件名", "状态", "备注"]


def export_report(rows: List[PlanRow], path: str = "") -> str:
    """导出对照表，返回生成的文件路径。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    if not path:
        os.makedirs(config.EXPORT_DIR, exist_ok=True)
        path = os.path.join(config.EXPORT_DIR,
                            "重命名对照表_%s.xlsx" % time.strftime("%Y%m%d_%H%M%S"))

    wb = Workbook()
    ws = wb.active
    ws.title = "重命名对照表"

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for c, h in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    warn_fill = PatternFill("solid", fgColor="FFE699")
    err_fill = PatternFill("solid", fgColor="F8CBAD")

    for i, r in enumerate(rows, start=2):
        values = [r.index, r.src_name, r.seq, r.name, r.new_name,
                  STATUS_TEXT.get(r.status, r.status), r.note]
        for c, v in enumerate(values, start=1):
            ws.cell(row=i, column=c, value=v)
        if r.status == "same":
            for c in range(1, len(_HEADERS) + 1):
                ws.cell(row=i, column=c).fill = warn_fill
        elif r.status not in ("ok",):
            for c in range(1, len(_HEADERS) + 1):
                ws.cell(row=i, column=c).fill = err_fill

    widths = [6, 45, 10, 14, 45, 12, 28]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
    ws.freeze_panes = "A2"

    wb.save(path)
    log.info("对照表已导出：%s（共 %d 行）", path, len(rows))
    return path